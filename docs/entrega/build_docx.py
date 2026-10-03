import base64
import copy
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
import time
import urllib.request
from pathlib import Path

import docx
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent
OUT_DIR = HERE / "salida"
DIAGRAM_DIR = HERE / "diagramas"
TEMPLATE = HERE / "plantilla_portada.docx"
SOURCES = ["01_arquitectura_y_documentacion_tecnica.md", "02_manual_de_usuario.md", "03_informe_de_pruebas.md"]

FONT = "Times New Roman"
BODY_PT = 12
TABLE_PT = 10.5
CODE_FONT = "Consolas"
CODE_PT = 9.5
FIRST_LINE_INDENT = Cm(1.27)
MARGIN = Cm(2.54)
LETTER = (Cm(21.59), Cm(27.94))
MAX_IMAGE_WIDTH = Cm(16)
PLACEHOLDER_HEIGHT = Cm(5.5)
CHROME = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
MERMAID_CDN = "https://cdn.jsdelivr.net/npm/mermaid@11/dist/mermaid.min.js"
DELIVERY_DATE = "25 de setiembre de 2026"

P_TIPO, P_TITULO, P_SUBTITULO, P_FECHA = 11, 12, 13, 27

def set_run_font(run, name=FONT, size=BODY_PT, bold=None, italic=None):
    run.font.name = name
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic

def para_format(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=FIRST_LINE_INDENT, spacing=1.5,
                before=0, after=0, left=None, hanging=None, keep_next=False):
    pf = p.paragraph_format
    p.alignment = align
    pf.first_line_indent = first_line
    if left is not None:
        pf.left_indent = left
    if hanging is not None:
        pf.first_line_indent = -hanging
    pf.line_spacing = spacing
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.keep_with_next = keep_next
    pf.widow_control = True

INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)")

def add_inline(p, text, size=BODY_PT, base_bold=False, base_italic=False):
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            add_inline(p, part[2:-2], size=size, base_bold=True, base_italic=base_italic)
        elif part.startswith("`") and part.endswith("`") and len(part) > 1:
            set_run_font(p.add_run(part[1:-1]), name=CODE_FONT, size=size - 1.5, bold=base_bold)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            set_run_font(p.add_run(part[1:-1]), size=size, bold=base_bold, italic=True)
        else:
            set_run_font(p.add_run(part), size=size, bold=base_bold, italic=base_italic)

def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)

def cell_borders(cell, **edges):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge, (val, sz, color) in edges.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), color)
        borders.append(el)

def fixed_width(table, widths_cm):
    tblpr = table._tbl.tblPr
    total = sum(widths_cm)
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(int(Cm(total).twips)))
    tblw.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(Cm(w).twips)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)

def table_no_borders(table):
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblpr.append(borders)

def paragraph_border(p, edge="bottom", sz=6, color="808080"):
    ppr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), "4")
    el.set(qn("w:color"), color)
    pbdr.append(el)
    ppr.append(pbdr)

def add_field(p, instr, placeholder=""):
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_begin)
    run2 = p.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    run2._r.append(it)
    run3 = p.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run3._r.append(sep)
    run4 = p.add_run(placeholder)
    set_run_font(run4)
    run5 = p.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run5._r.append(end)
    for r in (run, run2, run3, run5):
        set_run_font(r)

def make_template(source_docx):
    d = docx.Document(source_docx)
    body = d.element.body
    keep = P_FECHA + 1
    paras_seen = 0
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        if child.tag == qn("w:p"):
            paras_seen += 1
            if paras_seen <= keep:
                continue
        body.remove(child)
    d.save(TEMPLATE)

def set_paragraph_text(p, text, bold=None):
    runs = p.runs
    if not runs:
        r = p.add_run(text)
    else:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        r = runs[0]
    set_run_font(r, bold=bold)

def setup_styles(d):
    st = d.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(BODY_PT)
    rpr = st.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(0)
    for name, align, italic in (("Heading 1", WD_ALIGN_PARAGRAPH.CENTER, False),
                                ("Heading 2", WD_ALIGN_PARAGRAPH.LEFT, False),
                                ("Heading 3", WD_ALIGN_PARAGRAPH.LEFT, True)):
        h = d.styles[name]
        h.font.name = FONT
        h.font.size = Pt(BODY_PT)
        h.font.bold = True
        h.font.italic = italic
        h.font.color.rgb = RGBColor(0, 0, 0)
        hr = h.element.get_or_add_rPr()
        hf = hr.find(qn("w:rFonts"))
        if hf is None:
            hf = OxmlElement("w:rFonts")
            hr.insert(0, hf)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            hf.set(qn(attr), FONT)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if hf.get(qn(attr)) is not None:
                del hf.attrib[qn(attr)]
        pf = h.paragraph_format
        pf.alignment = align
        pf.page_break_before = False
        pf.first_line_indent = Cm(0)
        pf.left_indent = Cm(0)
        pf.line_spacing = 1.5
        pf.space_before = Pt(12)
        pf.space_after = Pt(0)
        pf.keep_with_next = True

def setup_page(d):
    for s in d.sections:
        s.page_width, s.page_height = LETTER
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = MARGIN
        s.different_first_page_header_footer = True
        header = s.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        for r in list(hp.runs):
            r._r.getparent().remove(r._r)
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_field(hp, "PAGE", "2")
        s.first_page_header.is_linked_to_previous = False
        for hf in (s.first_page_header, s.footer, s.first_page_footer):
            hf.is_linked_to_previous = False
            for fp in hf.paragraphs:
                for r in list(fp.runs):
                    r._r.getparent().remove(r._r)
                for fld in fp._p.findall(qn("w:fldSimple")):
                    fp._p.remove(fld)

def fill_cover(d, meta):
    ps = d.paragraphs
    set_paragraph_text(ps[P_TIPO], meta.get("tipo", "I Proyecto Programado"))
    set_paragraph_text(ps[P_TITULO], meta["titulo"], bold=True)
    set_paragraph_text(ps[P_SUBTITULO], meta.get("subtitulo", ""))
    set_paragraph_text(ps[P_FECHA], meta.get("fecha", DELIVERY_DATE))
    for p in ps[:P_FECHA + 1]:
        p.paragraph_format.first_line_indent = Cm(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            set_run_font(r, bold=r.bold)

class MermaidRenderer:
    def __init__(self):
        self.proc = None
        self.ws = None
        self.n = 0

    def _start(self):
        import websocket
        profile = tempfile.mkdtemp(prefix="mc-mermaid-")
        self.proc = subprocess.Popen(
            [CHROME, "--headless=new", "--disable-gpu", "--no-sandbox", "--hide-scrollbars",
             f"--user-data-dir={profile}", "--remote-debugging-port=9334", "--remote-allow-origins=*",
             "--window-size=1800,1400", "about:blank"],
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        for _ in range(60):
            try:
                tabs = json.load(urllib.request.urlopen("http://127.0.0.1:9334/json"))
                break
            except Exception:
                time.sleep(0.3)
        tab = [t for t in tabs if t.get("type") == "page"][0]
        self.ws = websocket.create_connection(tab["webSocketDebuggerUrl"], timeout=90)
        self.call("Page.enable")
        html = ("<html><head><meta charset='utf-8'><script src='%s'></script></head>"
                "<body style='margin:0;background:#fff'><div id='out'></div></body></html>") % MERMAID_CDN
        self.call("Page.navigate", url="data:text/html;base64," + base64.b64encode(html.encode()).decode())
        for _ in range(100):
            if self.js("typeof mermaid !== 'undefined'"):
                break
            time.sleep(0.2)
        self.js("mermaid.initialize({startOnLoad:false, theme:'neutral', securityLevel:'loose',"
                "fontFamily:'Times New Roman', themeVariables:{fontSize:'17px'},"
                "flowchart:{htmlLabels:true, curve:'basis'}})")

    def call(self, method, **params):
        self.n += 1
        self.ws.send(json.dumps({"id": self.n, "method": method, "params": params}))
        while True:
            r = json.loads(self.ws.recv())
            if r.get("id") == self.n:
                if "error" in r:
                    raise RuntimeError(f"{method}: {r['error']}")
                return r.get("result", {})

    def js(self, expr):
        r = self.call("Runtime.evaluate", expression=expr, returnByValue=True, awaitPromise=True)
        if r.get("exceptionDetails"):
            raise RuntimeError(r["exceptionDetails"].get("exception", {}).get("description", "error JS"))
        return r.get("result", {}).get("value")

    def render(self, code, out_png):
        if self.ws is None:
            self._start()
        self.js("document.getElementById('out').innerHTML=''")
        box = self.js(
            "(async () => { const {svg} = await mermaid.render('g%d', %s);"
            " const o=document.getElementById('out'); o.innerHTML=svg;"
            " const s=o.querySelector('svg'); s.style.maxWidth='none';"
            " const r=s.getBoundingClientRect(); return [r.x,r.y,r.width,r.height]; })()"
            % (self.n, json.dumps(code)))
        x, y, w, h = box
        self.call("Emulation.setDeviceMetricsOverride", width=int(x + w) + 40, height=int(y + h) + 40,
                  deviceScaleFactor=1, mobile=False)
        data = self.call("Page.captureScreenshot", format="png", captureBeyondViewport=True,
                         clip={"x": x, "y": y, "width": w, "height": h, "scale": 2})["data"]
        out_png.write_bytes(base64.b64decode(data))
        return w, h

    def close(self):
        if self.proc:
            self.proc.kill()

def parse_front_matter(text):
    meta = {}
    if text.startswith("---"):
        end = text.index("\n---", 3)
        for line in text[3:end].strip().splitlines():
            k, _, v = line.partition(":")
            meta[k.strip()] = v.strip()
        text = text[end + 4:]
    return meta, text

def build_document(src_name, renderer):
    src = HERE / src_name
    meta, text = parse_front_matter(src.read_text(encoding="utf-8"))
    d = docx.Document(TEMPLATE)
    setup_styles(d)
    setup_page(d)
    fill_cover(d, meta)

    body = d.element.body
    sectpr = body[-1]

    def new_par():
        p = d.add_paragraph()
        body.remove(p._p)
        sectpr.addprevious(p._p)
        return p

    def page_break():
        p = new_par()
        p.add_run().add_break(WD_BREAK.PAGE)
        para_format(p, first_line=Cm(0))

    def move_table(t):
        body.remove(t._tbl)
        sectpr.addprevious(t._tbl)

    page_break()
    h = new_par()
    h.style = d.styles["TOC Heading"] if "TOC Heading" in [s.name for s in d.styles] else d.styles["Heading 1"]
    h.text = ""
    set_run_font(h.add_run("Índice"), bold=True)
    para_format(h, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=Cm(0))
    toc = new_par()
    para_format(toc, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=Cm(0))
    add_field(toc, 'TOC \\o "1-2" \\h \\z \\u', "Actualice el índice: clic derecho → Actualizar campo.")
    page_break()

    lines = text.splitlines()
    i = 0
    stem = src.stem
    diagram_no = 0

    def caption(kind_line, above=True):
        m = re.match(r"(Tabla|Figura)\s+(\d+)\.?\s*(.*)", kind_line.strip())
        if not m:
            return
        p1 = new_par()
        set_run_font(p1.add_run(f"{m.group(1)} {m.group(2)}"), bold=True)
        para_format(p1, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=Cm(0), before=6, keep_next=True)
        p2 = new_par()
        add_inline(p2, m.group(3), base_italic=True)
        para_format(p2, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=Cm(0), after=4, keep_next=True)

    def add_picture(png):
        from PIL import Image
        with Image.open(png) as im:
            wpx, hpx = im.size
        width = min(MAX_IMAGE_WIDTH, Cm(wpx / 2 * 0.0264583))
        max_h = Cm(19)
        if width * hpx / wpx > max_h:
            width = int(max_h * wpx / hpx)
        p = new_par()
        para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=Cm(0), spacing=1.0, after=6)
        p.add_run().add_picture(str(png), width=width)

    pending_table_caption = None
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "\\pagebreak":
            page_break()
            i += 1
            continue

        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            level = len(m.group(1))
            p = new_par()
            p.style = d.styles[f"Heading {level}"]
            add_inline(p, m.group(2).strip(), base_bold=True, base_italic=(level == 3))
            i += 1
            continue

        if re.match(r"^(Tabla)\s+\d+\.", stripped) and i + 2 < len(lines) and (
                lines[i + 1].strip().startswith("|") or (not lines[i + 1].strip() and lines[i + 2].strip().startswith("|"))):
            pending_table_caption = stripped
            i += 1
            continue

        if stripped.startswith(":::figura") or stripped.startswith(":::diagrama"):
            kind = "figura" if stripped.startswith(":::figura") else "diagrama"
            cap = stripped.split(None, 1)[1] if len(stripped.split(None, 1)) > 1 else ""
            block = []
            i += 1
            while i < len(lines) and lines[i].strip() != ":::":
                block.append(lines[i])
                i += 1
            i += 1
            caption(cap)
            if kind == "diagrama":
                code = "\n".join(block)
                code = re.sub(r"^\s*```mermaid\s*\n|\n\s*```\s*$", "", code.strip())
                diagram_no += 1
                digest = hashlib.sha1(code.encode("utf-8")).hexdigest()[:10]
                png = DIAGRAM_DIR / f"{stem}_{diagram_no:02d}_{digest}.png"
                lucid = DIAGRAM_DIR / "lucid" / f"{stem}_{diagram_no:02d}.png"
                if lucid.exists():
                    png = lucid
                elif not png.exists():
                    for old in DIAGRAM_DIR.glob(f"{stem}_{diagram_no:02d}_*.png"):
                        old.unlink()
                    renderer.render(code, png)
                add_picture(png)
                continue
            images = [REPO / b.split(":", 1)[1].strip() for b in block if b.strip().startswith("Imagen:")]
            notes = [b.split(":", 1)[1].strip() for b in block if b.strip().startswith("Nota:")]
            if images and all(img.exists() for img in images):
                for img in images:
                    add_picture(img)
                for note in notes:
                    q = new_par()
                    set_run_font(q.add_run("Nota. "), size=10.5, italic=True)
                    add_inline(q, note, size=10.5)
                    para_format(q, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=Cm(0), spacing=1.15, after=10)
            else:
                t = d.add_table(rows=1, cols=1)
                move_table(t)
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                fixed_width(t, [16.5])
                cell = t.rows[0].cells[0]
                shade(cell, "F2F2F2")
                edge = ("dashed", 8, "7F7F7F")
                cell_borders(cell, top=edge, bottom=edge, left=edge, right=edge)
                trpr = t.rows[0]._tr.get_or_add_trPr()
                hgt = OxmlElement("w:trHeight")
                hgt.set(qn("w:val"), str(int(PLACEHOLDER_HEIGHT.twips)))
                hgt.set(qn("w:hRule"), "atLeast")
                trpr.append(hgt)
                cp = cell.paragraphs[0]
                set_run_font(cp.add_run("[Espacio para la imagen]"), size=10.5, bold=True)
                para_format(cp, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=Cm(0), spacing=1.15, before=6, after=4)
                for bl in [b for b in block if b.strip()]:
                    q = cell.add_paragraph()
                    add_inline(q, bl.strip(), size=10.5, base_italic=True)
                    para_format(q, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=Cm(0), spacing=1.15, after=3)
                spacer = new_par()
                para_format(spacer, first_line=Cm(0), spacing=1.0, after=6)
            continue

        if stripped.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            t = d.add_table(rows=1, cols=1)
            move_table(t)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            fixed_width(t, [16.5])
            cell = t.rows[0].cells[0]
            shade(cell, "F3F3F3")
            edge = ("single", 4, "BFBFBF")
            cell_borders(cell, top=edge, bottom=edge, left=edge, right=edge)
            first = True
            for bl in block or [""]:
                q = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                set_run_font(q.add_run(bl.rstrip()), name=CODE_FONT, size=CODE_PT)
                para_format(q, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=Cm(0), spacing=1.0)
            spacer = new_par()
            para_format(spacer, first_line=Cm(0), spacing=1.0, after=4)
            continue

        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if pending_table_caption:
                caption(pending_table_caption)
                pending_table_caption = None
            ncols = max(len(r) for r in rows)
            t = d.add_table(rows=len(rows), cols=ncols)
            move_table(t)
            table_no_borders(t)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, r in enumerate(rows):
                for ci in range(ncols):
                    cell = t.rows[ri].cells[ci]
                    txt = r[ci] if ci < len(r) else ""
                    cp = cell.paragraphs[0]
                    add_inline(cp, txt.replace("<br>", " "), size=TABLE_PT, base_bold=(ri == 0))
                    para_format(cp, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=Cm(0), spacing=1.0, before=2, after=2)
                    if ri == 0:
                        cell_borders(cell, top=("single", 8, "000000"), bottom=("single", 6, "000000"))
                        trpr = t.rows[0]._tr.get_or_add_trPr()
                        if trpr.find(qn("w:tblHeader")) is None:
                            trpr.append(OxmlElement("w:tblHeader"))
                    if ri == len(rows) - 1:
                        cell_borders(cell, bottom=("single", 8, "000000"))
            spacer = new_par()
            para_format(spacer, first_line=Cm(0), spacing=1.0, after=6)
            continue

        if stripped.startswith("> "):
            block = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                block.append(lines[i].strip()[1:].strip())
                i += 1
            p = new_par()
            add_inline(p, " ".join(block), size=11)
            para_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=Cm(0), left=Cm(1.0), spacing=1.15, before=4, after=6)
            ppr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            el = OxmlElement("w:left")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "18")
            el.set(qn("w:space"), "8")
            el.set(qn("w:color"), "7F7F7F")
            pbdr.append(el)
            ppr.append(pbdr)
            continue

        lm = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", line)
        if lm:
            while i < len(lines):
                lm = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", lines[i])
                if not lm:
                    if lines[i].strip() and lines[i].startswith("   ") and d.paragraphs:
                        add_inline(d.paragraphs[-1], " " + lines[i].strip())
                        i += 1
                        continue
                    break
                depth = 1 if len(lm.group(1)) >= 2 else 0
                marker = "•" if lm.group(2) in "-*" else lm.group(2)
                p = new_par()
                left = Cm(1.27 + 0.9 * depth)
                para_format(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=Cm(0), left=left, hanging=Cm(0.63))
                set_run_font(p.add_run(marker + "\t"))
                tabs = p.paragraph_format.tab_stops
                tabs.add_tab_stop(left)
                add_inline(p, lm.group(3).strip())
                i += 1
            continue

        block = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith(("#", "|", "```", ":::", "> ", "\\pagebreak"))
                    or re.match(r"^([-*]|\d+\.)\s+", nxt) or re.match(r"^Tabla\s+\d+\.", nxt)):
                break
            block.append(nxt)
            i += 1
        p = new_par()
        para_format(p)
        add_inline(p, " ".join(block))

    OUT_DIR.mkdir(exist_ok=True)
    name = meta.get("archivo") or f"{src.stem}.docx"
    out = OUT_DIR / name
    d.save(out)
    return out

def word_finalize(docx_paths):
    ps = "$ErrorActionPreference='Stop'; $w = New-Object -ComObject Word.Application; $w.Visible=$false; try {"
    for pth in docx_paths:
        pdf = str(pth.with_suffix(".pdf"))
        ps += (f" $d = $w.Documents.Open('{pth}'); $d.Fields.Update() | Out-Null;"
               f" foreach ($t in $d.TablesOfContents) {{ $t.Update() }};"
               f" $d.Save(); $d.ExportAsFixedFormat('{pdf}', 17); $d.Close();")
    ps += " } finally { $w.Quit() }"
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)

def main(argv):
    DIAGRAM_DIR.mkdir(exist_ok=True)
    if not TEMPLATE.exists():
        make_template(HERE.parents[3] / "Indagatoria_Gestion_Procesos_ArchLinux_Jon_Jen_Leno.docx")
    wanted = argv or SOURCES
    renderer = MermaidRenderer()
    outs = []
    try:
        for name in wanted:
            if (HERE / name).exists():
                outs.append(build_document(name, renderer))
                print("generado", outs[-1].name)
            else:
                print("falta la fuente", name)
    finally:
        renderer.close()
    if outs:
        word_finalize(outs)
        print("índice actualizado y PDF exportado con Word")

if __name__ == "__main__":
    main(sys.argv[1:])
